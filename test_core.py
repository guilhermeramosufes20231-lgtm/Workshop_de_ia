import io
import zipfile

from app.demo_data import demo_csv
from app.services.analysis import analyze
from app.services.csv_import import import_csv, parse_money
from app.services.excel_export import build_workbook
from app.services.simulator import simulate


def test_parse_money():
    assert parse_money("R$ 1.234,56") == (1234.56, False)
    assert parse_money("-45,90") == (45.9, True)
    assert parse_money("(100,00)") == (100.0, True)
    assert parse_money("1234.56") == (1234.56, False)
    assert parse_money("1.234.567,00") == (1234567.0, False)


def test_import_demo_personal():
    res = import_csv(demo_csv("pessoal"))
    r = res["report"]
    assert r["invalid_count"] == 2
    assert r["duplicates"] == 1
    assert r["card_double_count_avoided"] is True
    assert r["transfers"] == 6
    txs = res["transactions"]
    assert all(not t["counted"] for t in txs if t["type"] == "transferencia")
    assert all(not t["counted"] for t in txs if t["type"] == "pagamento_fatura")


def test_analysis_totals_match_manual_sum():
    txs = import_csv(demo_csv("pessoal"))["transactions"]
    a = analyze(txs, "pessoal")
    inc = sum(t["amount"] for t in txs if t["counted"] and t["type"] == "receita")
    exp = sum(t["amount"] for t in txs if t["counted"] and t["type"] in ("despesa", "pagamento_fatura"))
    assert abs(a["summary"]["income"] - inc) < 0.01
    assert abs(a["summary"]["expense"] - exp) < 0.01
    assert abs(sum(c["value"] for c in a["categories"]) - exp) < 0.05
    assert a["period"]["months"] == 6
    assert any(r["description"] == "Netflix" for r in a["recurring"])


def test_period_filter_and_no_data():
    txs = import_csv(demo_csv("pessoal"))["transactions"]
    a = analyze(txs, "pessoal", "2026-09-01", "2026-09-30")
    assert a["period"]["months"] == 1 and len(a["timeline"]) == 1
    assert analyze(txs, "pessoal", "2030-01-01", "2030-02-01")["has_data"] is False


def test_business_profile():
    a = analyze(import_csv(demo_csv("empresa"))["transactions"], "empresa")
    assert a["summary"]["costs"] > 0 and a["summary"]["gross_margin_pct"] is not None


def test_simulation():
    a = analyze(import_csv(demo_csv("pessoal"))["transactions"], "pessoal")
    top = a["categories"][0]
    s = simulate(a, {top["category"]: 10}, 5)
    assert abs(s["rows"][0]["saving_monthly"] - top["value"] / 6 * 0.1) < 0.02
    assert s["comparison"]["gain_monthly"] > s["comparison"]["saving_monthly"]


def test_excel_has_native_charts():
    txs = import_csv(demo_csv("pessoal"))["transactions"]
    recs = [{"tipo": "fato", "titulo": "t", "detalhe": "d", "premissas": "p"}]
    xlsx = build_workbook("pessoal", txs, "2026-04-01", "2026-09-30", recs, {"reductions": {"Alimentação": 15}, "revenue_increase_pct": 5})
    z = zipfile.ZipFile(io.BytesIO(xlsx))
    charts = [n for n in z.namelist() if n.startswith("xl/charts/chart")]
    assert len(charts) >= 5  # pizza, barras, colunas, linha, simulação
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames[:5] == ["Movimentações", "Resumo", "Categorias", "Análise temporal", "Recomendações"]
    assert wb["Movimentações"].max_row == len(txs) + 1
